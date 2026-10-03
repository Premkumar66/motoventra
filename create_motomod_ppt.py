import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Initialize Presentation with 16:9 Widescreen dimensions
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Color Palette
DARK_BG = RGBColor(6, 11, 24)       # #060b18
CARD_BG = RGBColor(12, 18, 34)      # #0c1222
CYAN = RGBColor(0, 242, 254)        # #00f2fe
PURPLE = RGBColor(155, 81, 224)     # #9b51e0
GREEN = RGBColor(16, 185, 129)      # #10b981
TEXT_WHITE = RGBColor(238, 242, 255)# #eef2ff
TEXT_MUTED = RGBColor(107, 127, 168)# #6b7fa8
ACCENT_AMBER = RGBColor(245, 158, 11)# #f59e0b

def set_slide_background(slide, color):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_header(slide, title_text, subtitle_text):
    # Top Accent Line
    top_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.1))
    top_line.fill.solid()
    top_line.fill.fore_color.rgb = CYAN
    top_line.line.fill.background()

    # Header Box
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.1))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p1 = tf.paragraphs[0]
    p1.text = title_text
    p1.font.name = "Arial"
    p1.font.size = Pt(26)
    p1.font.bold = True
    p1.font.color.rgb = CYAN

    p2 = tf.add_paragraph()
    p2.text = subtitle_text
    p2.font.name = "Arial"
    p2.font.size = Pt(13)
    p2.font.color.rgb = TEXT_MUTED

# ==============================================================================
# TITLE SLIDE (Cover Page)
# ==============================================================================
slide_layout = prs.slide_layouts[6] # Blank
slide_title = prs.slides.add_slide(slide_layout)
set_slide_background(slide_title, DARK_BG)

# Title Top Bar Accent
bar = slide_title.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.15))
bar.fill.solid()
bar.fill.fore_color.rgb = CYAN
bar.line.fill.background()

# Title Box
tb_title = slide_title.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(11.333), Inches(2.2))
tf = tb_title.text_frame
tf.word_wrap = True

p_main = tf.paragraphs[0]
p_main.text = "MOTOMOD AI"
p_main.font.name = "Arial"
p_main.font.size = Pt(44)
p_main.font.bold = True
p_main.font.color.rgb = CYAN

p_sub = tf.add_paragraph()
p_sub.text = "AI-Powered Motorcycle Modification & Performance Prediction Platform"
p_sub.font.name = "Arial"
p_sub.font.size = Pt(20)
p_sub.font.color.rgb = TEXT_WHITE
p_sub.font.bold = True

p_desc = tf.add_paragraph()
p_desc.text = "Department of Information Technology — Capstone Project Presentation"
p_desc.font.name = "Arial"
p_desc.font.size = Pt(14)
p_desc.font.color.rgb = TEXT_MUTED

# Presented By Card
card_authors = slide_title.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(4.2), Inches(11.333), Inches(2.3))
card_authors.fill.solid()
card_authors.fill.fore_color.rgb = CARD_BG
card_authors.line.color.rgb = PURPLE

tf_authors = card_authors.text_frame
tf_authors.word_wrap = True
tf_authors.vertical_anchor = MSO_ANCHOR.MIDDLE

p_hdr = tf_authors.paragraphs[0]
p_hdr.text = "PRESENTED BY:"
p_hdr.font.name = "Arial"
p_hdr.font.size = Pt(14)
p_hdr.font.bold = True
p_hdr.font.color.rgb = ACCENT_AMBER

p_a1 = tf_authors.add_paragraph()
p_a1.text = "1. PREMKUMAR R  (Reg No: 24IT0119)"
p_a1.font.name = "Arial"
p_a1.font.size = Pt(18)
p_a1.font.bold = True
p_a1.font.color.rgb = TEXT_WHITE

p_a2 = tf_authors.add_paragraph()
p_a2.text = "2. NANDEESH Y G  (Reg No: 24IT0103)"
p_a2.font.name = "Arial"
p_a2.font.size = Pt(18)
p_a2.font.bold = True
p_a2.font.color.rgb = TEXT_WHITE


# ==============================================================================
# SLIDE 1: PROBLEM IDENTIFICATION AND TITLE IDENTIFICATION
# ==============================================================================
slide1 = prs.slides.add_slide(slide_layout)
set_slide_background(slide1, DARK_BG)
add_header(slide1, "SLIDE 1: PROBLEM IDENTIFICATION & TITLE IDENTIFICATION", "Defining the core engineering problem and system title for MotoMod AI")

# Left Box: Title Identification
card_t = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.2))
card_t.fill.solid()
card_t.fill.fore_color.rgb = CARD_BG
card_t.line.color.rgb = CYAN

tf_t = card_t.text_frame
tf_t.word_wrap = True
p = tf_t.paragraphs[0]
p.text = "📌 TITLE IDENTIFICATION"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = CYAN

points_t = [
    ("Project Name", "MotoMod AI (Motoventra)"),
    ("Domain", "Automotive Informatics, Artificial Intelligence & Full-Stack Web"),
    ("Core Innovation", "ML-driven performance forecasting (BHP, Torque, Mileage, Top Speed) integrated with dynamic fitment verification"),
    ("Target Users", "Motorcycle enthusiasts, tuners, workshops, and aftermarket retailers"),
    ("Tech Stack", "FastAPI, Python ML (XGBoost/LightGBM), SQLite, Interactive HTML5/3D Visualizer")
]
for title, val in points_t:
    p = tf_t.add_paragraph()
    p.text = f"• {title}: {val}"
    p.font.size = Pt(13)
    p.font.color.rgb = TEXT_WHITE

# Right Box: Problem Identification
card_p = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.7), Inches(5.7), Inches(5.2))
card_p.fill.solid()
card_p.fill.fore_color.rgb = CARD_BG
card_p.line.color.rgb = PURPLE

tf_p = card_p.text_frame
tf_p.word_wrap = True
p = tf_p.paragraphs[0]
p.text = "⚠️ PROBLEM IDENTIFICATION"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = PURPLE

probs = [
    "1. Trial-and-Error Modifications: Riders install aftermarket parts without knowing real performance impacts or HP gains prior to purchase.",
    "2. Severe Incompatibility Risks: Mismatched components lead to engine damage, voided warranties, and safety hazards.",
    "3. Absence of Predictive Models: Traditional auto platforms only serve static part listings without dynamic performance forecasting.",
    "4. Scattered Specifications: Vehicle specs, fitment matrices, and review ratings are fragmented across incompatible platforms."
]
for pr in probs:
    p = tf_p.add_paragraph()
    p.text = pr
    p.font.size = Pt(13)
    p.font.color.rgb = TEXT_WHITE


# ==============================================================================
# SLIDE 2: LITERATURE SURVEY
# ==============================================================================
slide2 = prs.slides.add_slide(slide_layout)
set_slide_background(slide2, DARK_BG)
add_header(slide2, "SLIDE 2: LITERATURE SURVEY", "Comparative analysis of existing automotive systems vs. Proposed MotoMod AI Platform")

# Table
table_shape = slide2.shapes.add_table(4, 4, Inches(0.8), Inches(1.7), Inches(11.7), Inches(5.2))
table = table_shape.table
table.columns[0].width = Inches(2.5)
table.columns[1].width = Inches(3.0)
table.columns[2].width = Inches(3.2)
table.columns[3].width = Inches(3.0)

headers = ["System / Approach", "Primary Focus & Features", "Existing Limitations", "MotoMod AI Solution"]
for col_idx, h_text in enumerate(headers):
    cell = table.cell(0, col_idx)
    cell.fill.solid()
    cell.fill.fore_color.rgb = PURPLE
    p = cell.text_frame.paragraphs[0]
    p.text = h_text
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE

rows_data = [
    ("Traditional Auto E-Commerce (RevZilla, BikeWale)", "Static part catalogs with simple dropdown filters by make and model.", "No performance prediction algorithms, no 3D visualization, no dynamic fitment scoring.", "ML Ensemble prediction of HP, Torque, Mileage & Top Speed."),
    ("Generic Tuning Calculators", "Basic mathematical formulas for gear ratio and displacement math.", "Lacks real vehicle dataset integration, no ML models, no aftermarket part database.", "Integrated dataset of 262+ bike models across 46 brands with OEM specs."),
    ("Manual Workshop Diagnostics", "Physical dyno testing and manual trial-and-error fitment check.", "High cost, time-consuming, potential risk of mechanical component failure.", "Instant AI simulation, verified compatibility matrix, and UPI payment checkout.")
]

for row_idx, r_data in enumerate(rows_data, start=1):
    for col_idx, val in enumerate(r_data):
        cell = table.cell(row_idx, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = CARD_BG
        p = cell.text_frame.paragraphs[0]
        p.text = val
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_WHITE if col_idx != 3 else CYAN


# ==============================================================================
# SLIDE 3: OBJECTIVES OF PROJECT
# ==============================================================================
slide3 = prs.slides.add_slide(slide_layout)
set_slide_background(slide3, DARK_BG)
add_header(slide3, "SLIDE 3: OBJECTIVES OF PROJECT", "Core goals and technical milestones of the MotoMod AI Platform")

objs = [
    ("🎯 Primary Goal", "To build an AI-powered motorcycle modification & performance prediction ecosystem with real-time fitment verification.", CYAN),
    ("🤖 1. AI Performance Forecasting", "Develop ML models (XGBoost/LightGBM) to accurately predict HP gains (+/- 0.8 BHP precision), Torque shifts, Mileage, and Top Speed offsets.", TEXT_WHITE),
    ("🏍️ 2. Comprehensive Motorcycle Catalog", "Construct a structured database covering 262+ motorcycle variants across 46 global brands with complete technical specs.", TEXT_WHITE),
    ("🔍 3. Automated Fitment Matching", "Engineered direct and universal compatibility logic to strictly eliminate part incompatibility risks.", TEXT_WHITE),
    ("📸 4. 3D & 360° Visualizer", "Incorporate interactive 3D model viewer and multi-angle dataset images for realistic component inspection.", TEXT_WHITE),
    ("💳 5. UPI QR Payment & Checkout", "Build an integrated checkout workflow featuring instant scannable UPI QR codes and bank reference verification.", GREEN)
]

for idx, (title, desc, color) in enumerate(objs):
    row = idx // 2
    col = idx % 2
    x = Inches(0.8 + col * 5.9)
    y = Inches(1.7 + row * 1.7)
    
    card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.6), Inches(1.5))
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.color.rgb = color
    
    tf = card.text_frame
    tf.word_wrap = True
    
    p1 = tf.paragraphs[0]
    p1.text = title
    p1.font.size = Pt(15)
    p1.font.bold = True
    p1.font.color.rgb = color
    
    p2 = tf.add_paragraph()
    p2.text = desc
    p2.font.size = Pt(12)
    p2.font.color.rgb = TEXT_WHITE


# ==============================================================================
# SLIDE 4: PROJECT PLANNING AND TIME SCHEDULE (12 WEEKS)
# ==============================================================================
slide4 = prs.slides.add_slide(slide_layout)
set_slide_background(slide4, DARK_BG)
add_header(slide4, "SLIDE 4: PROJECT PLANNING AND TIME SCHEDULE (12 WEEKS)", "12-Week Gantt Chart Breakdown of Implementation Phases")

weeks_schedule = [
    ("Weeks 1 - 2", "Problem Identification, Literature Survey & Documentation (PRD / SRS)", PURPLE),
    ("Weeks 3 - 4", "Database Schema Design & Data Collection (262 Bikes, 46 Brands, Mod Parts)", CYAN),
    ("Weeks 5 - 6", "Machine Learning Model Development & Training (XGBoost / LightGBM)", GREEN),
    ("Weeks 7 - 8", "FastAPI Backend Architecture, Async Endpoints & Auth Implementation", ACCENT_AMBER),
    ("Weeks 9 - 10", "Frontend SPA, 3D / 360° Visualizer & Instant UPI QR Payment Gateway", CYAN),
    ("Weeks 11 - 12", "Integration Testing, Bug Fixes, System Verification & Final Presentation", GREEN)
]

for idx, (w_title, w_desc, bar_color) in enumerate(weeks_schedule):
    y = Inches(1.7 + idx * 0.85)
    
    # Label Box
    card_lbl = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y, Inches(2.4), Inches(0.7))
    card_lbl.fill.solid()
    card_lbl.fill.fore_color.rgb = CARD_BG
    card_lbl.line.color.rgb = bar_color
    tf = card_lbl.text_frame
    p = tf.paragraphs[0]
    p.text = w_title
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = bar_color
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    
    # Task Content Box
    card_task = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.4), y, Inches(9.1), Inches(0.7))
    card_task.fill.solid()
    card_task.fill.fore_color.rgb = CARD_BG
    card_task.line.color.rgb = bar_color
    tf_task = card_task.text_frame
    tf_task.word_wrap = True
    tf_task.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_task = tf_task.paragraphs[0]
    p_task.text = w_desc
    p_task.font.size = Pt(13)
    p_task.font.color.rgb = TEXT_WHITE


# ==============================================================================
# SLIDE 5: REFERENCE
# ==============================================================================
slide5 = prs.slides.add_slide(slide_layout)
set_slide_background(slide5, DARK_BG)
add_header(slide5, "SLIDE 5: REFERENCES", "Academic papers, technical documentation, and frameworks referenced")

refs = [
    ("1. Machine Learning for Vehicle Performance Estimation", "IEEE Transactions on Vehicular Technology, Ensemble Learning Algorithms (XGBoost & LightGBM) for Automotive Torque and Power Prediction."),
    ("2. FastAPI & Asynchronous Python Frameworks", "FastAPI Official Documentation (2024), Asynchronous Server Gateway Interface (ASGI) for High-Throughput RESTful APIs."),
    ("3. Relational Database Design for E-Commerce Fitment", "SQLAlchemy 2.0 & PostgreSQL Manual, Normalized Schema Architecture for Automotive Component Compatibility Mapping."),
    ("4. WebXR & Interactive 3D Model Rendering", "Google Model-Viewer & WebGL Standards, Real-Time 3D Object Rendering in Modern Web Browsers."),
    ("5. Unified Payments Interface (UPI) Integration Standards", "National Payments Corporation of India (NPCI) Technical Specification for Dynamic QR Code Generation and VPA Handshaking.")
]

for idx, (ref_title, ref_desc) in enumerate(refs):
    y = Inches(1.7 + idx * 1.0)
    card = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y, Inches(11.7), Inches(0.85))
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.color.rgb = CYAN
    
    tf = card.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    
    p1 = tf.paragraphs[0]
    p1.text = ref_title
    p1.font.size = Pt(14)
    p1.font.bold = True
    p1.font.color.rgb = CYAN
    
    p2 = tf.add_paragraph()
    p2.text = ref_desc
    p2.font.size = Pt(11)
    p2.font.color.rgb = TEXT_MUTED

# Save presentation
output_path = r"c:\CCP PROJECT\Motoventra\MotoMod_AI_Presentation.pptx"
prs.save(output_path)
print(f"SUCCESS: Created presentation at {output_path}")
